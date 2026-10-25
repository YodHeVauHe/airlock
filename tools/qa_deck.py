"""Programmatic QA for the deck: text overflow estimates + out-of-bounds + overlaps."""
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
prs = Presentation("submission/airlock-slides.pptx")
W = prs.slide_width / EMU_IN
H = prs.slide_height / EMU_IN

# rough per-font width factors (avg glyph width as fraction of font size, in inches/pt)
FACTOR = {"Consolas": 0.0075, "Segoe UI": 0.0060}

issues = 0
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        x, y = sh.left / EMU_IN, sh.top / EMU_IN
        w, h = sh.width / EMU_IN, sh.height / EMU_IN
        if x < -0.01 or y < -0.01 or x + w > W + 0.01 or y + h > H + 0.01:
            print(f"[s{si}] OUT OF BOUNDS: '{sh.text_frame.text[:40]}' box=({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
            issues += 1
        # estimated wrapped height
        est_h = 0.0
        for p in sh.text_frame.paragraphs:
            if not p.runs:
                continue
            size = max((r.font.size.pt if r.font.size else 18) for r in p.runs)
            font_name = p.runs[0].font.name or "Segoe UI"
            fac = FACTOR.get(font_name, 0.0062)
            line_w = sum(len(r.text) * (r.font.size.pt if r.font.size else 18) *
                         FACTOR.get(r.font.name or "Segoe UI", 0.0062) for r in p.runs)
            import math
            nlines = max(1, math.ceil(line_w / max(w, 0.1)))
            est_h += nlines * size * 1.28 / 72
            if p.space_after is not None:
                est_h += p.space_after.pt / 72
        if est_h > h * 1.35 and est_h - h > 0.25:
            print(f"[s{si}] OVERFLOW? est {est_h:.2f}in vs box {h:.2f}in: '{sh.text_frame.text[:48]}'")
            issues += 1
        if sh.text_frame.text.strip():
            boxes.append((x, y, w, h, sh.text_frame.text[:28]))
    # pairwise overlap of text boxes
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]
            bx, by, bw, bh, bt = boxes[j]
            ox = max(0, min(ax + aw, bx + bw) - max(ax, bx))
            oy = max(0, min(ay + ah, by + bh) - max(ay, by))
            if ox > 0.12 and oy > 0.12:
                print(f"[s{si}] OVERLAP: '{at}' x '{bt}' ({ox:.2f}x{oy:.2f}in)")
                issues += 1

print(f"\n{len(prs.slides._sldIdLst)} slides checked, {issues} issue(s)")
