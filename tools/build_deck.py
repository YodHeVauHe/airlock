"""Build the Airlock pitch deck (submission asset).

Visual identity mirrors the product UI: near-black mission-control background,
Airlock blue primary, containment green as the single accent, terminal-box motif.
Canvas: 13.33 x 7.5 in. Fonts: Segoe UI (titles/body) + Consolas (terminal).
"""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

# palette (the product's own)
BG = RGBColor(0x0D, 0x11, 0x17)       # near-black surface
PANEL = RGBColor(0x16, 0x1C, 0x28)    # slightly raised panel
PANEL2 = RGBColor(0x11, 0x16, 0x1F)   # darker panel
PRIMARY = RGBColor(0x58, 0xA6, 0xFF)  # airlock blue
ACCENT = RGBColor(0x3F, 0xB9, 0x50)   # containment green
TEXT = RGBColor(0xE6, 0xED, 0xF3)
MUTED = RGBColor(0x8B, 0x94, 0x9E)
DIM = RGBColor(0x6E, 0x76, 0x81)
RISK = RGBColor(0xF8, 0x51, 0x49)

SANS = "Segoe UI"
MONO = "Consolas"

W, H = 13.333, 7.5
M = 0.55

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, size=18, color=TEXT, bold=False, font=SANS,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [[(runs, {})]]          # [[(text, overrides)]] paragraphs of runs
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        if spacing is not None:
            p.space_after = Pt(spacing)
        for t, o in para:
            r = p.add_run(); r.text = t
            r.font.name = o.get("font", font)
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def terminal(s, x, y, w, h, lines, title="airlock ~ sandbox", size=12.5):
    """The deck's signature motif: a console window."""
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    box.adjustments[0] = 0.05
    box.fill.solid(); box.fill.fore_color.rgb = PANEL2
    box.line.color.rgb = RGBColor(0x1F, 0x27, 0x35); box.line.width = Pt(1)
    box.shadow.inherit = False
    # title bar dots + label
    for i, c in enumerate((RISK, RGBColor(0xD2, 0x99, 0x22), ACCENT)):
        d = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.16 + i * 0.22), Inches(y + 0.14), Inches(0.1), Inches(0.1))
        d.fill.solid(); d.fill.fore_color.rgb = c; d.line.fill.background(); d.shadow.inherit = False
    text(s, x + 0.9, y + 0.10, w - 1.1, 0.25, title, size=10.5, color=DIM, font=MONO)
    body = []
    for t, c in lines:
        body.append([(t, {"color": c})])
    text(s, x + 0.25, y + 0.45, w - 0.5, h - 0.6, body, size=size, font=MONO, spacing=4)
    return box


def chip(s, x, y, w, label, color=PRIMARY, filled=False):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    c.adjustments[0] = 0.5
    c.fill.solid() if filled else c.fill.background()
    if filled:
        c.fill.fore_color.rgb = color
    c.line.color.rgb = color; c.line.width = Pt(1.1)
    c.shadow.inherit = False
    tf = c.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = Emu(0); tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.name = MONO; r.font.size = Pt(11); r.font.bold = True
    r.font.color.rgb = BG if filled else color
    return c


def kicker(s, label, y=0.52):
    text(s, M, y, W - 2 * M, 0.3, label.upper(), size=13, color=PRIMARY, bold=True)


def title(s, t, y=0.86, size=32, w=W - 2 * M):
    text(s, M, y, w, 0.62, t, size=size, bold=True)


def footer(s, n):
    text(s, M, H - 0.42, 7, 0.3, "AIRLOCK  ·  blast-radius-zero agents on Vultr",
         size=10, color=DIM, font=MONO)
    text(s, W - M - 0.6, H - 0.42, 0.6, 0.3, f"{n:02d}", size=10, color=DIM,
         font=MONO, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 1 · title
s = slide()
kicker(s, "Vultr: Agent Rush Hackathon  ·  Challenge 1 — Blast Radius Zero", y=0.6)
text(s, M, 1.35, W - 2 * M, 1.5, [[("air", {"color": TEXT}), ("lock", {"color": PRIMARY})]],
     size=76, bold=True)
text(s, M, 2.85, 8.2, 1.1,
     "Give your AI agent real work.\nKeep your blast radius at zero.", size=25, color=MUTED)
chip(s, M, 4.35, 1.5, "plan", PRIMARY, filled=True)
chip(s, M + 1.7, 4.35, 1.7, "dispatch", PRIMARY, filled=True)
chip(s, M + 3.6, 4.35, 1.5, "verify", PRIMARY, filled=True)
chip(s, M + 5.3, 4.35, 1.6, "destroy", ACCENT, filled=True)
terminal(s, M, 5.15, 7.6, 1.75, [
    ("$ rm -rf /", TEXT),
    ("sandbox destroyed — host untouched        CONTAINED", ACCENT),
    ("$ receipt sealed: 27 events · sha256 chain intact", MUTED),
])
text(s, 8.6, 5.15, W - M - 8.6, 1.8,
     [[("github.com/YodHeVauHe/airlock", {"color": PRIMARY, "font": MONO, "size": 14})],
      [("FastAPI · Docker · Vultr VM", {"color": DIM, "size": 12.5})],
      [("Vultr Serverless Inference", {"color": DIM, "size": 12.5})],
      [("18 tests · CI green · MIT", {"color": DIM, "size": 12.5})]], spacing=6)

# ------------------------------------------------- 2 · the problem
s = slide()
kicker(s, "The problem")
title(s, "One hallucination away from your machine")
terminal(s, M, 2.0, 5.7, 3.1, [
    ("# a typical agent host", MUTED),
    ("$ export AIRLOCK_INFERENCE_API_KEY=vlt-…", TEXT),
    ("$ python agent.py \"clean up temp files\"", TEXT),
    ("", TEXT),
    ("rm -rf /          # model wrote this", RISK),
    ("FATAL: host filesystem erased", RISK),
], title="someone-else-laptop ~ agent", size=13)
rows = [
    ("Your secrets", "API keys sit in the same env the model can read"),
    ("Your network", "nothing stops a call home, or a prompt-injection fetch"),
    ("Your data", "one wrong path deletes production files, not test data"),
]
y = 2.0
for t, d in rows:
    text(s, 6.75, y, 6.0, 0.8, [[(t, {"bold": True, "size": 19})], [(d, {"color": MUTED, "size": 14})]], spacing=3)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.75), Inches(y + 0.83), Inches(6.0), Pt(1))
    ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0x1F, 0x27, 0x35); ln.line.fill.background(); ln.shadow.inherit = False
    y += 1.05
text(s, 6.75, y + 0.1, 6.0, 0.9,
     [[("An agent that only chats is a demo.", {"color": TEXT, "bold": True, "size": 17})],
      [("An agent that executes safely is a product.", {"color": PRIMARY, "bold": True, "size": 17})]], spacing=4)
footer(s, 2)

# ------------------------------------------------- 3 · what we built
s = slide()
kicker(s, "What we built")
title(s, "Airlock: real agent work, executed inside a disposable cell")
text(s, M, 1.75, 11.8, 0.6,
     "A production-style web app on a Vultr VM. You hand the agent real work; Airlock plans it with "
     "Vultr Serverless Inference, executes generated code in a per-task sandbox, streams every action "
     "live — and destroys the sandbox afterwards.", size=15.5, color=MUTED)
lanes = [
    ("Data analyst", "upload a CSV, get computed answers,\nsummary tables and rendered charts"),
    ("File fixer", "upload a broken file, get a repaired\nversion and a report of what changed"),
    ("Chart anything", "describe the visual, get a real\nrendered artifact from executed code"),
]
x = M
for t, d in lanes:
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.75), Inches(3.85), Inches(1.75))
    card.adjustments[0] = 0.06
    card.fill.solid(); card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = RGBColor(0x1F, 0x27, 0x35); card.line.width = Pt(1); card.shadow.inherit = False
    text(s, x + 0.25, 2.95, 3.35, 0.4, t, size=17, bold=True, color=PRIMARY)
    text(s, x + 0.25, 3.42, 3.35, 0.95, d.replace("\n", " "), size=13, color=MUTED)
    x += 4.02
text(s, M, 4.95, 12.2, 0.5, [[("Every action is ", {"size": 18, "color": TEXT}),
     ("streamed live, risk-scored, and sealed into a tamper-evident receipt.", {"size": 18, "color": ACCENT, "bold": True})]])
terminal(s, M, 5.6, W - 2 * M, 1.35, [
    ("task 7f54db  ·  policy verdict: clean  ·  sandbox net:none mem:512MB pids:128  ·  exit 0  ·  artifacts: summary.csv chart.svg", MUTED),
    ("receipt: sha256 a3f9…c41e ← 9d02…77b1 ← 5c88…09d3 …  VERIFIED ✓", ACCENT),
], title="airlock ~ task timeline", size=13)
footer(s, 3)

# ------------------------------------------------- 4 · architecture
s = slide()
kicker(s, "Architecture")
title(s, "One control plane, disposable cells, verifiable receipts")
# browser
b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(M), Inches(2.1), Inches(2.3), Inches(1.0))
b.adjustments[0] = 0.1; b.fill.solid(); b.fill.fore_color.rgb = PANEL
b.line.color.rgb = PRIMARY; b.line.width = Pt(1.2); b.shadow.inherit = False
text(s, M, 2.32, 2.3, 0.6, "Browser", size=15, bold=True, align=PP_ALIGN.CENTER)
# control plane
cp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.0), Inches(1.75), Inches(3.6), Inches(2.9))
cp.adjustments[0] = 0.05; cp.fill.solid(); cp.fill.fore_color.rgb = PANEL
cp.line.color.rgb = PRIMARY; cp.line.width = Pt(1.4); cp.shadow.inherit = False
text(s, 4.2, 1.9, 3.2, 0.35, "Control plane — Vultr VM", size=14, bold=True, color=PRIMARY)
for i, (t, d) in enumerate([
    ("Planner", "plan → code → fix → explain"),
    ("Policy engine", "pre-exec scan + risk score"),
    ("Orchestrator", "lifecycle: create → run → destroy"),
    ("Receipts", "sha256 hash chain + seal"),
]):
    y = 2.35 + i * 0.56
    text(s, 4.25, y, 3.2, 0.5, [[(t, {"size": 13.5, "bold": True})], [(d, {"size": 11, "color": MUTED})]], spacing=1)
# vultr inference
vi = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(1.75), Inches(3.45), Inches(1.15))
vi.adjustments[0] = 0.08; vi.fill.solid(); vi.fill.fore_color.rgb = PANEL
vi.line.color.rgb = RGBColor(0x1F, 0x27, 0x35); vi.line.width = Pt(1); vi.shadow.inherit = False
text(s, 9.5, 1.95, 3.1, 0.8, [[("Vultr Serverless Inference", {"size": 14, "bold": True})],
                              [("https://api.vultrinference.com/v1", {"size": 10.5, "color": MUTED, "font": MONO})]], spacing=3)
# sandbox drivers
sd = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(3.5), Inches(3.45), Inches(1.85))
sd.adjustments[0] = 0.06; sd.fill.solid(); sd.fill.fore_color.rgb = PANEL
sd.line.color.rgb = ACCENT; sd.line.width = Pt(1.4); sd.shadow.inherit = False
text(s, 9.5, 3.62, 3.1, 0.35, "Sandbox — one per task", size=14, bold=True, color=ACCENT)
text(s, 9.5, 4.05, 3.1, 1.25, [
    [("docker", {"font": MONO, "size": 12, "bold": True}), ("  net:none · caps:0 · ro-rootfs", {"size": 11.5, "color": MUTED})],
    [("local", {"font": MONO, "size": 12, "bold": True}), ("  dev driver (labeled)", {"size": 11.5, "color": MUTED})],
    [("vultr", {"font": MONO, "size": 12, "bold": True}), ("  throwaway VM per task", {"size": 11.5, "color": MUTED})],
], spacing=5)
# connectors
def arrow(x1, y1, x2, y2, color=PRIMARY):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(1.6)
    tail = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    he = tail.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    tail.append(he)
arrow(M + 2.3, 2.6, 3.98, 2.6)              # browser -> control plane (SSE)
arrow(7.6, 2.3, 9.28, 2.3)                  # planner -> inference
arrow(7.6, 4.4, 9.28, 4.4, color=ACCENT)    # orchestrator -> sandbox
text(s, 2.88, 2.28, 1.05, 0.25, "HTTP/SSE", size=9.5, color=DIM, font=MONO)
text(s, 7.75, 2.02, 1.5, 0.25, "HTTPS", size=9.5, color=DIM, font=MONO)
text(s, 7.75, 4.12, 1.5, 0.25, "create/kill", size=9.5, color=DIM, font=MONO)
terminal(s, M, 5.55, W - 2 * M, 1.4, [
    ("user goal -> plan (Vultr Inference) -> generated code -> policy gate -> fresh sandbox -> real output -> receipt sealed -> sandbox destroyed", MUTED),
], title="airlock ~ happy path", size=13)
footer(s, 4)

# ------------------------------------------------- 5 · invariants
s = slide()
kicker(s, "Containment-first, enforced")
title(s, "Four invariants the product never bends")
inv = [
    ("Process isolation", "agent code never runs in the control plane — only inside a disposable boundary, as an unprivileged user with zero capabilities"),
    ("Secret hygiene", "sandbox environments are built from scratch — API keys never enter a cell, and policy blocks env probes before execution"),
    ("Resource limits", "wall-clock kill, memory / CPU / PID ceilings, output caps — enforced by the driver, never by the model's good behavior"),
    ("Lifecycle discipline", "every sandbox is destroyed after its task, on every path — timeout, crash, success — and the destroy is a logged event"),
]
y = 1.95
for i, (t, d) in enumerate(inv, start=1):
    num = text(s, M, y - 0.06, 0.75, 0.6, [[(f"0{i}", {"size": 30, "bold": True, "color": PRIMARY, "font": MONO})]])
    text(s, M + 0.85, y, 11.3, 0.75,
         [[(t + "   ", {"size": 18, "bold": True}), (d, {"size": 14.5, "color": MUTED})]])
    if i < 4:
        ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M + 0.85), Inches(y + 0.78), Inches(11.3), Pt(1))
        ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0x1F, 0x27, 0x35); ln.line.fill.background(); ln.shadow.inherit = False
    y += 1.13
terminal(s, M, 6.35, W - 2 * M, 0.72, [
    ("challenge asks for all four invariants — Airlock enforces them and proves them on demand", MUTED),
], size=12.5)
footer(s, 5)

# ------------------------------------------------- 6 · torture chamber
s = slide()
kicker(s, "The containment moment — built into the product")
title(s, "The Torture Chamber: don't take containment on faith")
text(s, M, 1.7, 12.2, 0.5,
     "One click detonates real attacks inside the sandbox — streamed and receipted like any task. "
     "A host fingerprint hash is taken before and after every run.", size=15, color=MUTED)
cards = [
    ("rm -rf /", "bounces off the read-only rootfs and the unprivileged user"),
    ("fork bomb", "the pids ceiling absorbs the fan-out; the run dies at the cap"),
    ("exfiltrate env", "the cell env is scratch-built — there is nothing to steal"),
    ("phone home", "no network namespace — every connect fails fast"),
]
x = M
for t, d in cards:
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.5), Inches(2.95), Inches(2.15))
    card.adjustments[0] = 0.07
    card.fill.solid(); card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = RGBColor(0x1F, 0x27, 0x35); card.line.width = Pt(1); card.shadow.inherit = False
    text(s, x + 0.2, 2.68, 2.55, 0.4, t, size=15.5, bold=True, color=RISK, font=MONO)
    text(s, x + 0.2, 3.15, 2.55, 1.35, d, size=12.5, color=MUTED)
    x += 3.12
text(s, M, 5.0, 12.2, 0.75, [[("CONTAINED", {"size": 34, "bold": True, "color": ACCENT, "font": MONO}),
     ("  —  host fingerprint unchanged, on every detonation", {"size": 17, "color": TEXT})]])
terminal(s, M, 5.85, W - 2 * M, 1.2, [
    ("host.proof before: f803ec53…  →  rm -rf / attempted  →  sandbox.destroyed  →  host.proof after: f803ec53…  ✓", MUTED),
], title="airlock ~ torture run 7f54db", size=13)
footer(s, 6)

# ------------------------------------------------- 7 · receipts
s = slide()
kicker(s, "Verifiable output")
title(s, "Every run ends with a receipt anyone can verify")
text(s, M, 1.75, 12.2, 0.55,
     "Events form a per-task SHA-256 hash chain. The sealed receipt — the chain plus a summary — "
     "downloads from the UI; one command recomputes it and detects any tampering.", size=15, color=MUTED)
# hash chain visual: 5 linked blocks
bx, bw, gap = M + 0.15, 1.85, 0.55
labels = ["plan.ready", "code.gen", "policy ✓", "exec exit 0", "sealed"]
for i, lab in enumerate(labels):
    blk = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx + i * (bw + gap)), Inches(2.7), Inches(bw), Inches(1.05))
    blk.adjustments[0] = 0.12
    blk.fill.solid(); blk.fill.fore_color.rgb = PANEL2
    blk.line.color.rgb = PRIMARY if i < 4 else ACCENT; blk.line.width = Pt(1.2); blk.shadow.inherit = False
    text(s, bx + i * (bw + gap), 2.82, bw, 0.3, lab, size=12, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    text(s, bx + i * (bw + gap), 3.18, bw, 0.4, f"{9700+i*731:04x}…{4100+i*977:04x}", size=10, color=MUTED, font=MONO, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(bx + i * (bw + gap) + bw + 0.04, 3.22, bx + (i + 1) * (bw + gap) - 0.04, 3.22, color=PRIMARY)
# links back visual
text(s, M, 3.95, 12.2, 0.35, "each block commits to the previous block's hash — remove one, and every later seal breaks",
     size=12.5, color=DIM, font=MONO)
pts = [
    ("Real, not vibes", "the explainer is handed the run's actual stdout and artifacts — charts exist because code wrote them"),
    ("Download & verify", "GET /api/tasks/{id}/receipt — verification is a few lines, no Airlock internals required"),
    ("Judge-friendly", "the demo video's receipt moment doubles as an integrity proof of the run you just watched"),
]
y = 4.6
for t, d in pts:
    text(s, M, y, 12.2, 0.55, [[(t + "   ", {"size": 16, "bold": True, "color": PRIMARY}), (d, {"size": 14, "color": MUTED})]])
    y += 0.72
footer(s, 7)

# ------------------------------------------------- 8 · built on vultr
s = slide()
kicker(s, "Built on Vultr — the challenge checklist")
title(s, "Vultr is the system of control, not just a host")
rows = [
    ("Vultr Cloud Compute VM", "runs the control plane, the artifact store — and the sandboxes themselves", "required ✓"),
    ("Vultr Serverless Inference", "every planner, fixer and explainer call — OpenAI-compatible endpoint, zero GPUs", "required ✓"),
    ("Vultr API — throwaway instances", "experimental driver: a whole disposable VM per task, created by cloud-init, destroyed in a finally", "stretch ✓"),
    ("Vultr credits", "the entire demo runs inside the event's $200 credit budget", "efficient ✓"),
]
y = 2.0
for t, d, tag in rows:
    text(s, M, y, 4.4, 0.4, t, size=16.5, bold=True)
    text(s, M + 4.55, y + 0.02, 6.1, 0.7, d, size=13.5, color=MUTED)
    chip(s, W - M - 1.35, y - 0.02, 1.35, tag, ACCENT if "✓" in tag else PRIMARY, filled=False)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(y + 0.72), Inches(W - 2 * M), Pt(1))
    ln.fill.solid(); ln.fill.fore_color.rgb = RGBColor(0x1F, 0x27, 0x35); ln.line.fill.background(); ln.shadow.inherit = False
    y += 1.0
terminal(s, M, 6.15, W - 2 * M, 0.9, [
    ("deploy: bash deploy/setup-vm.sh  →  Docker installed, sandbox image built, control plane live on port 80  (~5 min)", MUTED),
], size=13)
footer(s, 8)

# ------------------------------------------------- 9 · business value
s = slide()
kicker(s, "Business value")
title(s, "The trust layer between LLM agents and production")
text(s, M, 1.8, 12.2, 0.55,
     "Enterprises want agentic automation on real data and real workflows. What they cannot accept "
     "is an untrusted model holding the keys. Airlock sells exactly that missing piece.", size=15, color=MUTED)
stats = [("3", "sandbox drivers,\none boundary contract"), ("18", "tests incl. containment\n+ receipt tamper-evidence"),
         ("4", "torture detonations,\nall absorbed live"), ("100", "community-standards\nscore on the repo")]
x = M
for v, lab in stats:
    text(s, x, 2.75, 2.9, 0.9, [[(v, {"size": 56, "bold": True, "color": PRIMARY, "font": MONO})]])
    text(s, x, 3.72, 2.9, 0.75, lab.replace("\n", " "), size=12, color=MUTED)
    x += 3.05
rows2 = [
    ("Data & analytics teams", "let an agent crunch sensitive exports without giving it your laptop, your S3 keys, or your network"),
    ("Platform / DevOps", "a reviewed, auditable pattern for untrusted code execution — receipts slot into existing compliance flows"),
    ("Support & ops automation", "agents that act, not just chat — with a hard ceiling on what a bad day can destroy"),
]
y = 4.75
for t, d in rows2:
    text(s, M, y, 12.2, 0.5, [[(t + "   ", {"size": 16, "bold": True}), (d, {"size": 13.5, "color": MUTED})]])
    y += 0.68
footer(s, 9)

# ------------------------------------------------- 10 · close
s = slide()
kicker(s, "Roadmap & close")
title(s, "From hackathon build to production default")
road = [
    ("Pattern B — browser lane", "Playwright inside the same cell contract; screenshots verified by a vision pass"),
    ("gVisor driver", "kernel-level interception for stronger process isolation"),
    ("Egress allowlist proxy", "deny-by-default today; allowlisted destinations tomorrow"),
    ("Multi-tenant workspaces", "teams, quotas and per-run ACLs on the same receipt rails"),
]
x = M
for t, d in road:
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(2.95), Inches(1.9))
    card.adjustments[0] = 0.07
    card.fill.solid(); card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = RGBColor(0x1F, 0x27, 0x35); card.line.width = Pt(1); card.shadow.inherit = False
    text(s, x + 0.2, 2.2, 2.55, 0.6, t, size=14, bold=True, color=PRIMARY)
    text(s, x + 0.2, 2.85, 2.55, 1.0, d, size=11.5, color=MUTED)
    x += 3.12
terminal(s, M, 4.35, W - 2 * M, 1.5, [
    ("$ airlock — plan → dispatch → verify → destroy", TEXT),
    ("real executed work · zero blast radius · signed proof", ACCENT),
    ("github.com/YodHeVauHe/airlock  ·  MIT  ·  built on Vultr", MUTED),
], title="airlock ~ closing", size=14)
text(s, M, 6.2, 12.2, 0.5, "Thank you — ask us to detonate something live.",
     size=16, color=TEXT, bold=True)
footer(s, 10)

prs.save("submission/airlock-slides.pptx")
print("saved submission/airlock-slides.pptx with", len(prs.slides._sldIdLst), "slides")
